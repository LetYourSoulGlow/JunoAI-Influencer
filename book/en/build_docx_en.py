"""Build an editable DOCX of 'The Silence I Come From' for a publisher (no ISBN — the publisher assigns it).

Usage: python3 build_docx_en.py <translation_en.json> <out.docx>
Same block rules as build_book_en.py: chapter openers, exercise headers, sub-headings, bullet lists.
"""
import sys, re, json
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
tr = json.load(open(src))
paras = [''] * (max(map(int, tr)) + 2)
for k, v in tr.items(): paras[int(k)] = v.strip()
END = re.compile(r'[.!?…:”"»),;—-]$')
FONT = 'Garamond'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(148), Mm(210)  # A5
sec.top_margin, sec.bottom_margin = Mm(18), Mm(20)
sec.left_margin, sec.right_margin = Mm(18), Mm(16)
sec.gutter = Mm(2)

def setfont(style, size, bold=False, italic=False, color=None):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    style.font.size = Pt(size); style.font.bold = bold; style.font.italic = italic
    if color: style.font.color.rgb = RGBColor(*color)

st = doc.styles
setfont(st['Normal'], 11)
pf = st['Normal'].paragraph_format
pf.space_after = Pt(0); pf.line_spacing = 1.2; pf.first_line_indent = Mm(5)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.widow_control = True
for name, size, bold, italic in [('Heading 1', 18, False, False), ('Heading 2', 13, True, False), ('Heading 3', 11.5, True, True)]:
    s = st[name]; setfont(s, size, bold, italic, (0x11, 0x11, 0x11))
    s.paragraph_format.first_line_indent = Mm(0); s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
st['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
st['Heading 1'].paragraph_format.space_before = Pt(6); st['Heading 1'].paragraph_format.space_after = Pt(24)
st['Heading 2'].paragraph_format.space_before = Pt(14); st['Heading 2'].paragraph_format.space_after = Pt(6)
st['Heading 3'].paragraph_format.space_before = Pt(10); st['Heading 3'].paragraph_format.space_after = Pt(4)
lb = st['List Bullet']; setfont(lb, 11); lb.paragraph_format.first_line_indent = Mm(0)
lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; lb.paragraph_format.space_after = Pt(2)

def para(text='', style=None, align=None, size=None, italic=False, bold=False, indent=None, before=0, after=0, caps=False, spacing=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if indent is not None: p.paragraph_format.first_line_indent = indent
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.italic = italic; r.bold = bold; r.font.all_caps = caps
        if size: r.font.size = Pt(size)
        if spacing is not None:
            rpr = r._r.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(spacing)); rpr.append(sp)
    return p

def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

C = WD_ALIGN_PARAGRAPH.CENTER
# Title page
para('Joanna Dorobisz', align=C, size=13, caps=True, spacing=40, indent=Mm(0), before=110)
para('The Silence', align=C, size=28, indent=Mm(0), before=60)
para('I Come From', align=C, size=28, indent=Mm(0), after=18)
para('Coming Back to Who You Truly Are', align=C, size=14, italic=True, indent=Mm(0))
page_break()
# Copyright page (no ISBN: assigned by the publisher)
for i in range(18): para(indent=Mm(0))
for t, it in [('Copyright © Joanna Dorobisz', False),
              ('Translated from the Polish Cisza, z której pochodzę.', False),
              ('All rights reserved. No part of this book may be reproduced or distributed without the author’s written permission, except for brief quotations in reviews.', False),
              ('This book is not a substitute for therapy or professional help. If you are experiencing abuse or are in crisis, please contact a professional or a local helpline.', False),
              ('Let Your Soul Glow · letyoursoulglow.store', False)]:
    p = para(align=WD_ALIGN_PARAGRAPH.LEFT, indent=Mm(0), after=6)
    if t.startswith('Translated'):
        p.add_run('Translated from the Polish ').font.size = Pt(9)
        r = p.add_run('Cisza, z której pochodzę'); r.italic = True; r.font.size = Pt(9)
        p.add_run('.').font.size = Pt(9)
    else:
        p.add_run(t).font.size = Pt(9)
page_break()
# Contents: static list of chapters (the publisher sets page numbers; headings use Heading 1 for an automatic TOC)
para('Contents', align=C, size=16, indent=Mm(0), after=14)
for i, t in enumerate(paras):
    m = re.match(r'^Chapter\s+(\d+):\s*(.+)$', t)
    if m:
        para(f'Chapter {m.group(1)}', size=8, caps=True, spacing=20, indent=Mm(0), align=WD_ALIGN_PARAGRAPH.LEFT)
        para(m.group(2), size=10, indent=Mm(0), align=WD_ALIGN_PARAGRAPH.LEFT, after=5)
    elif t.startswith('Conclusion:'):
        para('Conclusion', size=8, caps=True, spacing=20, indent=Mm(0), align=WD_ALIGN_PARAGRAPH.LEFT)
        para(t.split(':', 1)[1].strip(), size=10, indent=Mm(0), align=WD_ALIGN_PARAGRAPH.LEFT, after=5)
    elif t == 'Epilogue':
        para('Epilogue', size=10, indent=Mm(0), align=WD_ALIGN_PARAGRAPH.LEFT, after=5)

# Body
first_after_heading = True
ul = []
def flush():
    global ul, first_after_heading
    for x in ul: para(x, style='List Bullet')
    if ul: first_after_heading = True
    ul = []

n_ch = 0
for i, t in enumerate(paras):
    if i == 0 or not t:
        if not t: flush()
        continue
    if t.startswith('•'):
        ul.append(t.lstrip('•').strip()); continue
    flush()
    m = re.match(r'^Chapter\s+(\d+):\s*(.+)$', t)
    if m or t.startswith('Conclusion:') or t == 'Epilogue':
        if m: label, title = f'Chapter {m.group(1)}', m.group(2)
        elif t == 'Epilogue': label, title = '', 'Epilogue'
        else: label, title = 'Conclusion', t.split(':', 1)[1].strip()
        page_break(); n_ch += 1
        for _ in range(4): para(indent=Mm(0))
        if label: para(label, align=C, size=10, caps=True, spacing=40, indent=Mm(0), after=6)
        doc.add_paragraph(title, style='Heading 1')
        first_after_heading = True; continue
    if t.startswith('Exercises and Reflective Questions'):
        p = doc.add_paragraph(t, style='Heading 2'); p.alignment = C; p.paragraph_format.space_before = Pt(20)
        first_after_heading = True; continue
    if i >= len(paras) - 4 and t == 'Joanna Dorobisz':
        para(t, align=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, indent=Mm(0), before=6); continue
    if re.match(r'^\d+\.\s', t) and len(t) < 100:
        doc.add_paragraph(t, style='Heading 3'); first_after_heading = True; continue
    if len(t) < 75 and not END.search(t):
        doc.add_paragraph(t, style='Heading 2'); first_after_heading = True; continue
    p = doc.add_paragraph(t)
    if first_after_heading: p.paragraph_format.first_line_indent = Mm(0)
    first_after_heading = False
flush()

# Page numbers in footer
fp = sec.footer.paragraphs[0]; fp.alignment = C
r = fp.add_run()
for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if typ:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ); r._r.append(e)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)
r.font.size = Pt(9)
doc.core_properties.title = 'The Silence I Come From'; doc.core_properties.author = 'Joanna Dorobisz'
doc.save(out)
print('chapters:', n_ch)
